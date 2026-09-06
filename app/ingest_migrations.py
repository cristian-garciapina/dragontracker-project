"""Farlight migrations Excel ingestion pipeline.

Source: cod-game-tools.farlightgames.com/migration (official beta portal).
Two sheets: "Arrivées" (IN to K544) and "Départs" (OUT from K544).

See migrations-tracker-spec.md for design decisions.
"""
from __future__ import annotations

import logging
import unicodedata
from datetime import date, datetime
from pathlib import Path
from io import BytesIO
from typing import Any, Union

import openpyxl
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from sqlalchemy.orm import Session

from .models import Migration

logger = logging.getLogger(__name__)


HEADER_MAP_COMMON = {
    "rang": "rank",
    "identifiant du personnage": "character_id",
    "nom du personnage": "name_at_migration",
    "puissance": "power_at_migration",
    "heure de migration": "migration_date",
    "score de migration": "migration_score",
}

# Sheet-specific header for the other kingdom column
HEADER_OTHER_KINGDOM = {
    "id. royaume source": "other_kingdom",   # Arrivées
    "id. royaume cible": "other_kingdom",    # Départs
}

SHEET_ARRIVEES = "Arrivées"
SHEET_DEPARTS = "Départs"


def _normalize(s: Any) -> str:
    """Lowercase, strip accents, collapse whitespace."""
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return " ".join(s.split())


def _parse_int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(str(v).strip())
    except (ValueError, TypeError):
        return None


def _parse_date(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None



def _open_workbook(source: Union[str, Path, bytes, BytesIO], **kw):
    """Load a workbook from a path, raw bytes, or a BytesIO."""
    if isinstance(source, (bytes, bytearray)):
        return openpyxl.load_workbook(BytesIO(source), **kw)
    return openpyxl.load_workbook(source, **kw)


def is_migrations_file(source: Union[str, Path, bytes, BytesIO]) -> bool:
    """Return True if the xlsx has 'Arrivées' + 'Départs' sheets."""
    try:
        wb = _open_workbook(source, read_only=True, data_only=True)
    except Exception as e:
        logger.warning("is_migrations_file: cannot open workbook: %s", e)
        return False
    names = set(wb.sheetnames)
    wb.close()
    return SHEET_ARRIVEES in names and SHEET_DEPARTS in names


def _parse_sheet(ws, direction: str) -> tuple[list[dict], int]:
    """Parse a single sheet. Returns (rows, skipped_count)."""
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], 0

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        norm = _normalize(cell)
        if norm in HEADER_MAP_COMMON:
            col_map[idx] = HEADER_MAP_COMMON[norm]
        elif norm in HEADER_OTHER_KINGDOM:
            col_map[idx] = HEADER_OTHER_KINGDOM[norm]

    required = {"character_id", "name_at_migration", "power_at_migration",
                "other_kingdom", "migration_date"}
    missing = required - set(col_map.values())
    if missing:
        raise ValueError(
            f"Sheet missing required columns: {missing}. Found: {set(col_map.values())}"
        )

    parsed: list[dict] = []
    skipped = 0
    for raw in rows_iter:
        if raw is None or all(v is None for v in raw):
            continue
        rec: dict[str, Any] = {}
        for idx, field in col_map.items():
            rec[field] = raw[idx] if idx < len(raw) else None

        cid = _parse_int(rec.get("character_id"))
        power = _parse_int(rec.get("power_at_migration"))
        other = _parse_int(rec.get("other_kingdom"))
        mdate = _parse_date(rec.get("migration_date"))
        name = rec.get("name_at_migration")
        score = _parse_int(rec.get("migration_score"))

        if cid is None or other is None or mdate is None or not name:
            skipped += 1
            continue

        parsed.append({
            "character_id": cid,
            "direction": direction,
            "other_kingdom": other,
            "migration_date": mdate,
            "migration_score": score,
            "power_at_migration": power,
            "name_at_migration": str(name)[:64],
        })
    return parsed, skipped


def parse_migrations_xlsx(source: Union[str, Path, bytes, BytesIO]) -> dict[str, Any]:
    """Parse both sheets. Returns dict with incoming/outgoing rows + skip counts.

    Does NOT touch the DB. Pure function, safe for dry-run.
    """
    wb = _open_workbook(source, data_only=True)
    incoming, skipped_in = _parse_sheet(wb[SHEET_ARRIVEES], "IN")
    outgoing, skipped_out = _parse_sheet(wb[SHEET_DEPARTS], "OUT")
    wb.close()
    return {
        "incoming": incoming,
        "outgoing": outgoing,
        "skipped_incoming": skipped_in,
        "skipped_outgoing": skipped_out,
    }


def ingest_migrations(
    source: Union[str, Path, bytes, BytesIO],
    session: Session,
    source_filename: str,
) -> dict[str, int]:
    """Parse the xlsx and insert into `migrations` table with idempotence.

    Idempotence: relies on uq_migration (character_id, direction,
    migration_date, other_kingdom). Uses SQLite ON CONFLICT DO NOTHING.
    """
    parsed = parse_migrations_xlsx(source)
    now = datetime.utcnow()

    counters = {
        "incoming_inserted": 0,
        "outgoing_inserted": 0,
        "incoming_skipped_parse": parsed["skipped_incoming"],
        "outgoing_skipped_parse": parsed["skipped_outgoing"],
        "incoming_duplicates": 0,
        "outgoing_duplicates": 0,
    }

    for direction_key, rows in (("incoming", parsed["incoming"]),
                                 ("outgoing", parsed["outgoing"])):
        for r in rows:
            payload = {
                **r,
                "source_filename": source_filename,
                "ingested_at": now,
            }
            stmt = (
                sqlite_insert(Migration)
                .values(**payload)
                .on_conflict_do_nothing(index_elements=[
                    "character_id", "direction", "migration_date", "other_kingdom"
                ])
            )
            result = session.execute(stmt)
            if result.rowcount == 1:
                counters[f"{direction_key}_inserted"] += 1
            else:
                counters[f"{direction_key}_duplicates"] += 1

    session.commit()
    logger.info("ingest_migrations(%s): %s", source_filename, counters)
    return counters


def ingest_api_migrations(
    session: Session,
    incoming: list[dict],
    outgoing: list[dict],
    source_filename: str,
) -> dict[str, int]:
    """Insert already-mapped API migration rows into the `migrations` table.

    `incoming` / `outgoing` must be produced by
    farlight_client.map_api_migration_rows (dicts with character_id,
    other_kingdom, migration_date [str or date], power_at_migration,
    name_at_migration, migration_score). `direction`, `source_filename`
    and `ingested_at` are added here.

    Idempotence: same ON CONFLICT DO NOTHING on uq_migration.
    """
    now = datetime.utcnow()
    counters = {
        "incoming_inserted": 0,
        "outgoing_inserted": 0,
        "incoming_duplicates": 0,
        "outgoing_duplicates": 0,
    }

    for direction_key, direction_flag, rows in (
        ("incoming", "IN", incoming),
        ("outgoing", "OUT", outgoing),
    ):
        for r in rows:
            mdate = r.get("migration_date")
            if isinstance(mdate, str):
                try:
                    mdate = datetime.strptime(mdate.strip(), "%Y-%m-%d").date()
                except ValueError:
                    continue
            if mdate is None:
                continue

            payload = {
                "character_id": r["character_id"],
                "direction": direction_flag,
                "other_kingdom": r["other_kingdom"],
                "migration_date": mdate,
                "migration_score": r.get("migration_score"),
                "power_at_migration": r.get("power_at_migration"),
                "name_at_migration": r["name_at_migration"],
                "source_filename": source_filename,
                "ingested_at": now,
            }
            stmt = (
                sqlite_insert(Migration)
                .values(**payload)
                .on_conflict_do_nothing(index_elements=[
                    "character_id", "direction", "migration_date", "other_kingdom"
                ])
            )
            result = session.execute(stmt)
            if result.rowcount == 1:
                counters[f"{direction_key}_inserted"] += 1
            else:
                counters[f"{direction_key}_duplicates"] += 1

    session.commit()
    logger.info("ingest_api_migrations(%s): %s", source_filename, counters)
    return counters
